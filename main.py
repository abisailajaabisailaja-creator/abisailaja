from fastapi import FastAPI, Request, Form
from fastapi.responses import HTMLResponse, Response, RedirectResponse
from fastapi.templating import Jinja2Templates

from reportlab.pdfgen import canvas
from docx import Document

from io import BytesIO

from ai_service import generate_legal_document
from database import create_database, save_document, get_documents


app = FastAPI()

templates = Jinja2Templates(directory="templates")

create_database()


@app.get("/", response_class=HTMLResponse)
async def login_page(request: Request):
    return templates.TemplateResponse(
        "login.html",
        {"request": request}
    )


@app.post("/login", response_class=HTMLResponse)
async def login(
    request: Request,
    username: str = Form(...),
    password: str = Form(...)
):
    if username == "admin" and password == "1234":
        return RedirectResponse(
            url="/home",
            status_code=303
        )

    return templates.TemplateResponse(
        "login.html",
        {
            "request": request,
            "error": "Invalid User ID or Password"
        }
    )


@app.get("/logout")
async def logout():
    return RedirectResponse(
        url="/",
        status_code=303
    )


@app.get("/home", response_class=HTMLResponse)
async def home(request: Request):
    return templates.TemplateResponse(
        "index.html",
        {"request": request}
    )


@app.get("/form", response_class=HTMLResponse)
async def form(
    request: Request,
    type: str = "Employment Contract"
):
    return templates.TemplateResponse(
        "form.html",
        {
            "request": request,
            "document_type": type
        }
    )


@app.post("/generate", response_class=HTMLResponse)
async def generate(
    request: Request,
    document_type: str = Form("Employment Contract"),
    name: str = Form(""),
    company: str = Form(""),
    date: str = Form(""),
    role: str = Form(""),
    address: str = Form(""),
    terms: str = Form("")
):
    try:
        document = generate_legal_document(
            document_type,
            name,
            company,
            date,
            role,
            address,
            terms
        )

        save_document(
            document_type,
            name,
            company,
            date,
            role,
            address,
            terms,
            document
        )

        return templates.TemplateResponse(
            "result.html",
            {
                "request": request,
                "document_type": document_type,
                "document": document
            }
        )

    except Exception as e:
        return HTMLResponse(
            f"""
            <html>
            <body style="font-family:Arial;text-align:center;padding:60px">
                <h1>⚖️ LegalEase</h1>
                <h2>Error: {str(e)}</h2>
                <a href="/home">Back to Home</a>
            </body>
            </html>
            """
        )


@app.get("/history", response_class=HTMLResponse)
async def history(request: Request):
    documents = get_documents()

    return templates.TemplateResponse(
        "history.html",
        {
            "request": request,
            "documents": documents
        }
    )


@app.post("/download-pdf")
async def download_pdf(document: str = Form(...)):

    buffer = BytesIO()

    pdf = canvas.Canvas(buffer)

    width = 595
    height = 842

    left_margin = 50
    right_margin = 50
    top_margin = 50
    bottom_margin = 50

    y = height - top_margin

    for line in document.split("\n"):

        line = line.strip()

        if not line:
            y -= 12
            continue

        if y < bottom_margin:
            pdf.showPage()
            y = height - top_margin

        if line.isupper() and len(line) < 80:

            pdf.setFont(
                "Helvetica-Bold",
                14
            )

            pdf.drawCentredString(
                width / 2,
                y,
                line[:80]
            )

            y -= 25

        elif (
            len(line) > 2
            and line[0].isdigit()
            and "." in line[:4]
        ):

            pdf.setFont(
                "Helvetica-Bold",
                11
            )

            pdf.drawString(
                left_margin,
                y,
                line[:95]
            )

            y -= 20

        else:

            pdf.setFont(
                "Helvetica",
                10
            )

            words = line.split()

            current_line = ""

            for word in words:

                test_line = current_line + " " + word

                if pdf.stringWidth(
                    test_line,
                    "Helvetica",
                    10
                ) > (
                    width
                    - left_margin
                    - right_margin
                ):

                    pdf.drawString(
                        left_margin,
                        y,
                        current_line.strip()
                    )

                    y -= 15

                    if y < bottom_margin:
                        pdf.showPage()
                        y = height - top_margin

                    current_line = word

                else:
                    current_line = test_line

            if current_line.strip():

                pdf.drawString(
                    left_margin,
                    y,
                    current_line.strip()
                )

                y -= 15

    pdf.save()

    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        media_type="application/pdf",
        headers={
            "Content-Disposition":
            "attachment; filename=LegalEase_Document.pdf"
        }
    )


@app.post("/download-docx")
async def download_docx(document: str = Form(...)):

    doc = Document()

    title = doc.add_heading(
        "LegalEase",
        0
    )

    title.alignment = 1

    for line in document.split("\n"):

        line = line.strip()

        if not line:
            doc.add_paragraph()
            continue

        if line.isupper() and len(line) < 80:

            doc.add_heading(
                line,
                level=1
            )

            continue

        elif (
            len(line) > 2
            and line[0].isdigit()
            and "." in line[:4]
        ):

            doc.add_heading(
                line,
                level=2
            )

            continue

        doc.add_paragraph(line)

    buffer = BytesIO()

    doc.save(buffer)

    buffer.seek(0)

    return Response(
        buffer.getvalue(),
        media_type=(
            "application/"
            "vnd.openxmlformats-officedocument."
            "wordprocessingml.document"
        ),
        headers={
            "Content-Disposition":
            "attachment; filename=LegalEase_Document.docx"
        }
    )
